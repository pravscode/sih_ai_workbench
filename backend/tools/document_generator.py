from docx import Document
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from pathlib import Path


OUTPUT_DIR = Path("outputs")
OUTPUT_DIR.mkdir(exist_ok=True)


def generate_word(title, content):
    file_path = OUTPUT_DIR / "generated_report.docx"

    document = Document()

    document.add_heading(title, level=1)

    for paragraph in content.split("\n"):
        if paragraph.strip():
            document.add_paragraph(paragraph)

    document.save(file_path)

    return str(file_path)


def generate_pdf(title, content):
    file_path = OUTPUT_DIR / "generated_report.pdf"

    document = SimpleDocTemplate(str(file_path))

    styles = getSampleStyleSheet()

    elements = []

    elements.append(
        Paragraph(title, styles["Title"])
    )

    elements.append(Spacer(1, 20))

    for paragraph in content.split("\n"):
        if paragraph.strip():
            elements.append(
                Paragraph(paragraph, styles["BodyText"])
            )
            elements.append(Spacer(1, 10))

    document.build(elements)

    return str(file_path)